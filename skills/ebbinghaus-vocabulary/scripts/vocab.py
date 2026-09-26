#!/usr/bin/env python3
"""Deterministic local vocabulary scheduling and print generation."""
from __future__ import annotations

import argparse
import csv
import io
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, timedelta
from html.parser import HTMLParser

OFFSETS = (0, 2, 6, 14, 29)
FIELDS = ('english', 'chinese', 'ipa', 'pos')
REQUIRED_FIELDS = ('english', 'chinese')
OPTIONAL_FIELDS = ('ipa', 'pos')
FIELD_LABELS = {'chinese': '中文', 'english': '英文', 'ipa': '音标', 'pos': '词性'}
ALIASES = {'english': 'english', 'word': 'english', '英文': 'english', '英文单词': 'english',
           'chinese': 'chinese', 'meaning': 'chinese', '中文': 'chinese', '中文释义': 'chinese',
           'ipa': 'ipa', '音标': 'ipa', 'pos': 'pos', '词性': 'pos'}
ROOT = Path(__file__).resolve().parents[1]
RUNTIME = Path.home() / '.cache/codex-runtimes/codex-primary-runtime/dependencies'
PDF_W, PDF_H = 841.8897637795, 595.2755905512
MARGIN, COL_PAD = 24, 10
COL_W = (PDF_W - 2 * MARGIN) / 4
CONTENT_W = COL_W - 2 * COL_PAD
CN_W = 60
BODY_H = PDF_H - 2 * MARGIN - 42
MASTER_SCHEMAS = {
    ('ipa', 'pos'): (['chinese', 'english', 'ipa', 'pos'], [100, 65, 75, 24]),
    ('ipa',): (['chinese', 'english', 'ipa'], [112, 72, 80]),
    ('pos',): (['chinese', 'english', 'pos'], [120, 114, 30]),
    (): (['chinese', 'english', 'chinese', 'english'], [55, 77, 55, 77]),
}
MASTER_COLUMN_SCALE = .93  # Reserve width for office font substitution at print time.
MASTER_ROWS = 126
MASTER_START = 5
FONT_NAME = None


def font_name():
    global FONT_NAME
    if FONT_NAME:
        return FONT_NAME
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    windows_fonts = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts'
    candidates = [os.environ.get('VOCAB_FONT', ''), '/Library/Fonts/Arial Unicode.ttf',
                  str(windows_fonts / 'msyh.ttc'), str(windows_fonts / 'simsun.ttc'),
                  '/usr/share/fonts/truetype/arphic/ukai.ttc']
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            pdfmetrics.registerFont(TTFont('VocabularyCJK', candidate))
            FONT_NAME = 'VocabularyCJK'
            return FONT_NAME
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    FONT_NAME = 'STSong-Light'
    return FONT_NAME


def wrap(text, width, size):
    """Measured wrapping with word boundaries and Chinese punctuation handling.

    Only whitespace at a newly introduced line boundary is discarded. Original
    text remains in the input records; no meaning is shortened for printing.
    """
    from reportlab.pdfbase.pdfmetrics import stringWidth
    if width <= 0:
        raise ValueError('换行宽度必须大于零。')
    closing = set('，。；：！？、）》】」』〕〉,.!?;:)]}')
    opening = set('（《【「『〔〈([{')
    lines = []
    measure = lambda value: stringWidth(value, font_name(), size)
    for paragraph in str(text).split('\n'):
        remaining = paragraph.strip()
        if not remaining:
            lines.append('')
        while remaining:
            end = 0
            while end < len(remaining) and measure(remaining[:end + 1]) <= width:
                end += 1
            if end == 0:
                raise ValueError('单个字符超出可打印宽度。')
            if end < len(remaining):
                # Prefer a space boundary when a Latin word would be broken.
                if remaining[end - 1].isascii() and remaining[end - 1].isalnum() and remaining[end].isascii() and remaining[end].isalnum():
                    space = remaining.rfind(' ', 0, end)
                    if space > 0:
                        end = space
                # Move characters to the next line, never grow beyond width.
                while end > 1 and (remaining[end] in closing or remaining[end - 1] in opening):
                    end -= 1
            lines.append(remaining[:end].rstrip())
            remaining = remaining[end:].lstrip()
    return lines or ['']


def clean(value):
    return unicodedata.normalize('NFC', '' if value is None else str(value)).strip()


def records_from_table(rows):
    rows = [list(row) for row in rows if any(clean(v) for v in row)]
    if not rows:
        raise ValueError('词表为空。')
    headers = [ALIASES.get(clean(v).lower(), '') for v in rows[0]]
    if not {'english', 'chinese'} <= set(headers):
        raise ValueError('需要包含英文和中文的表头；请先用 Skill 整理材料。')
    if len([h for h in headers if h]) != len(set(h for h in headers if h)):
        raise ValueError('存在重复字段表头。')
    return [{h: row[i] if i < len(row) else '' for i, h in enumerate(headers) if h}
            for row in rows[1:]]


def load_words(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == '.json':
        raw = json.loads(path.read_text(encoding='utf-8-sig'))
        raw = raw.get('words', []) if isinstance(raw, dict) else raw
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        populated = [s for s in wb if any(any(v is not None for v in r) for r in s.values)]
        if len(populated) != 1:
            wb.close()
            raise ValueError('多工作表Excel请先提取、复核并合并词条，不自动漏掉其他表。')
        raw = records_from_table(populated[0].values)
        wb.close()
    elif suffix in ('.csv', '.tsv', '.txt'):
        text = path.read_text(encoding='utf-8-sig')
        raw = records_from_table(csv.reader(io.StringIO(text), delimiter=',' if suffix == '.csv' else '\t'))
    else:
        raise ValueError('build接受JSON/CSV/TSV/带表头文本/XLSX；其他材料先extract并复核。')
    return normalize(raw)


def normalize(raw):
    if not isinstance(raw, list) or not raw:
        raise ValueError('词表必须是非空记录数组。')
    words, seen, variants = [], set(), {}
    duplicates = 0
    for index, record in enumerate(raw, 1):
        if not isinstance(record, dict):
            raise ValueError(f'第{index}条不是记录对象。')
        if record.get('needs_review'):
            raise ValueError(f'第{index}条待复核，不能生成正式练习。')
        word = {k: clean(record.get(k)) for k in REQUIRED_FIELDS}
        # Absence is represented by an omitted optional key, never invented content.
        word.update({k: clean(record[k]) for k in OPTIONAL_FIELDS if clean(record.get(k))})
        if not word['english'] or not word['chinese']:
            raise ValueError(f'第{index}条缺少英文或中文，请先核对。')
        if any(re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', v) for v in word.values()):
            raise ValueError(f'第{index}条包含不可打印控制字符。')
        key = tuple(word.get(k, '') for k in FIELDS)
        if key in seen:
            duplicates += 1
            continue
        seen.add(key)
        words.append(word)
        variants[word['english']] = variants.get(word['english'], 0) + 1
    warnings = {'exact_duplicates_removed': duplicates,
                'same_word_variants': [k for k, v in variants.items() if v > 1],
                'missing_ipa': sum(not w.get('ipa') for w in words),
                'missing_pos': sum(not w.get('pos') for w in words)}
    return words, warnings


def schedule(words, daily, start=None):
    if daily not in (5, 10):
        raise ValueError('每天新增数只能为5或10。')
    if not words:
        raise ValueError('没有词条。')
    groups = [{'group': i // daily + 1, 'words': words[i:i + daily]}
              for i in range(0, len(words), daily)]
    days = {d: [] for d in range(1, len(groups) + 30)}
    for group in groups:
        for visit, offset in enumerate(OFFSETS, 1):
            days[group['group'] + offset].append({'group': group['group'], 'visit': visit,
                                                 'words': group['words']})
    start_date = date.fromisoformat(start) if start else None
    result = []
    for day, tasks in days.items():
        tasks.sort(key=lambda t: t['visit'])
        result.append({'day': day, 'date': (start_date + timedelta(days=day - 1)).isoformat() if start else None,
                       'tasks': tasks, 'word_count': sum(len(t['words']) for t in tasks)})
    peak = max(d['word_count'] for d in result)
    return {'word_count': len(words), 'daily': daily, 'group_count': len(groups),
            'total_days': len(groups) + 29, 'peak_words': peak,
            'peak_days': [d['day'] for d in result if d['word_count'] == peak],
            'theoretical_peak': daily * 5, 'start': start,
            'end': result[-1]['date'], 'groups': groups, 'days': result}


def task_layout(task, size=10, row_min=22):
    rows = []
    for word in task['words']:
        lines = wrap(word['chinese'], CN_W - 5, size)
        stacked = len(lines) > 2
        if stacked:
            lines = wrap(word['chinese'], CONTENT_W - 8, size)
        height = len(lines) * (size + 2) + (24 if stacked else 6)
        rows.append({'lines': lines, 'stacked': stacked, 'height': max(row_min, height)})
    return {**task, 'rows': rows, 'height': 20 + sum(r['height'] for r in rows) + 8,
            'font_size': size}


def day_columns(day):
    for size, row_min in ((10, 22), (10, 20), (9, 20)):
        columns, current, used = [], [], 0
        for task in day['tasks']:
            block = task_layout(task, size, row_min)
            if block['height'] > BODY_H:
                break
            if current and used + block['height'] > BODY_H:
                columns.append(current)
                current, used = [], 0
            current.append(block)
            used += block['height']
        else:
            if current:
                columns.append(current)
            if len(columns) <= 4:
                return [{'day': day['day'], 'date': day['date'], 'blocks': c} for c in columns]
    raise ValueError(f"Day {day['day']}的完整组在四列内放不下。请核对并简化过长中文释义后重试。")


def pdf_layout(plan):
    pages, current = [], []
    for day in plan['days']:
        if not day['tasks']:
            continue
        columns = day_columns(day)
        if len(current) + len(columns) > 4:
            pages.append(current)
            current = []
        current.extend(columns)
    if current:
        pages.append(current)
    return pages


def master_layout(groups):
    blocks = []
    for group in groups:
        optional = tuple(k for k in OPTIONAL_FIELDS if any(w.get(k) for w in group['words']))
        fields, widths = MASTER_SCHEMAS[optional]
        paired = not optional
        count = len(group['words'])
        half = (count + 1) // 2
        # Read the left pair top to bottom, then the right pair, preserving source order.
        index_rows = [(i, i + half if i + half < count else None) for i in range(half)] if paired else [(i,) for i in range(count)]
        rows = []
        for indexes in index_rows:
            if paired:
                values = [group['words'][idx].get(k, '') if idx is not None else ''
                          for idx in indexes for k in ('chinese', 'english')]
            else:
                values = [group['words'][indexes[0]].get(k, '') for k in fields]
            # A little width reserve for Excel's different font metrics.
            lines = [wrap(value, (width - 12) * .94, 9) for value, width in zip(values, widths)]
            units = max(3, math.ceil((max(map(len, lines)) * 12 + 8) / 6))
            rows.append({'values': ['\n'.join(s) for s in lines], 'units': units,
                         'word_indexes': list(indexes)})
        units = 6 + sum(r['units'] for r in rows) + 1
        if units > MASTER_ROWS - MASTER_START + 1:
            raise ValueError(f"第{group['group']}组太高，无法完整放入Excel半页。请缩短过长字段。")
        blocks.append({'group': group['group'], 'rows': rows, 'units': units,
                       'fields': fields, 'headers': [FIELD_LABELS[k] for k in fields],
                       'widths': widths, 'mode': 'paired' if paired else 'standard'})
    pages, page, side, cursor = [], [], 0, MASTER_START
    for block in blocks:
        if cursor + block['units'] - 1 > MASTER_ROWS:
            side += 1
            cursor = MASTER_START
        if side > 1:
            pages.append(page)
            page, side = [], 0
        page.append({**block, 'side': side, 'row': cursor})
        cursor += block['units']
    if page:
        pages.append(page)
    return pages


def master_grid(layout):
    """One physical sheet grid supports differently shaped groups via merged cells."""
    edges = {0}
    for page in layout:
        for block in page:
            position = 0
            for width in block['widths']:
                position += width
                edges.add(position)
            assert position == 264
    edges = sorted(edges)
    widths = [b - a for a, b in zip(edges, edges[1:])]
    return {'edges': edges, 'area_columns': len(widths),
            'widths': widths + [14] + widths,
            'last_column': excel_column(len(widths) * 2 + 1)}


def block_columns(block, grid):
    position, columns = 0, []
    for width in block['widths']:
        left = grid['edges'].index(position)
        position += width
        columns.append((left, grid['edges'].index(position) - left))
    return columns


def excel_column(number):
    letters = ''
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def compare_plans(words):
    return [{'daily': daily, **{key: plan[key] for key in ('word_count', 'group_count', 'total_days', 'peak_words')}}
            for daily in (5, 10) for plan in [schedule(words, daily)]]


def render_pdf(pages, path):
    from reportlab.pdfgen.canvas import Canvas
    c = Canvas(str(path), pagesize=(PDF_W, PDF_H), pageCompression=1)
    c.setTitle('艾宾浩斯复习打印册')
    c.setAuthor('艾宾浩斯式单词复习')
    font = font_name()
    for columns in pages:
        # Cut/fold guides run on the physical column boundaries.
        c.setStrokeColorRGB(.72, .72, .72)
        c.setLineWidth(.35)
        c.setDash(2, 4)
        for divider in range(1, 4):
            x = MARGIN + divider * COL_W
            c.line(x, MARGIN, x, PDF_H - MARGIN)
        c.setDash()
        for i, column in enumerate(columns):
            x = MARGIN + i * COL_W + COL_PAD
            y = PDF_H - MARGIN
            c.setFillColorRGB(.08, .08, .08)
            c.setFont(font, 13)
            c.drawString(x, y - 14, f"Day {column['day']}")
            c.setFont(font, 9)
            c.drawString(x, y - 31, '日期：____________')
            y -= 42
            for block in column['blocks']:
                c.setFillColorRGB(.94, .94, .94)
                c.rect(x, y - 18, CONTENT_W, 18, stroke=0, fill=1)
                c.setFillColorRGB(.12, .12, .12)
                c.setFont(font, 10)
                c.drawString(x + 4, y - 13, f"第{block['group']}组｜第{block['visit']}次")
                y -= 20
                c.setFont(font, block['font_size'])
                for row in block['rows']:
                    size = block['font_size']
                    line_height = size + 2
                    baseline = (y - size - 3 if row['stacked'] else
                                y - (row['height'] - len(row['lines']) * line_height) / 2 - size)
                    for j, line in enumerate(row['lines']):
                        c.drawString(x + 2, baseline - j * line_height, line)
                    c.setStrokeColorRGB(.40, .40, .40)
                    c.setLineWidth(.45)
                    line_start = x + 2 if row['stacked'] else x + CN_W + 3
                    c.line(line_start, y - row['height'] + 5, x + CONTENT_W - 2, y - row['height'] + 5)
                    y -= row['height']
                y -= 8
            if y < MARGIN - .01:
                raise AssertionError('PDF内容超出页面。')
        c.showPage()
    c.save()


def patch_print_settings(path, page_count, grid):
    """Add missing native print features after artifact-tool XLSX export."""
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    ET.register_namespace('', ns)
    q = lambda s: f'{{{ns}}}{s}'
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    sheet = ET.fromstring(files['xl/worksheets/sheet1.xml'])
    # Excel column units depend on the Normal style's maximum digit width.
    # Use Arial 10 (7 pixels at 96 dpi) and identical final widths for both
    # exporters. Their own pixel-to-column conversions are not interchangeable.
    styles = ET.fromstring(files['xl/styles.xml'])
    fonts = styles.find(q('fonts'))
    normal_font = ET.SubElement(fonts, q('font'))
    ET.SubElement(normal_font, q('sz'), {'val': '10'})
    ET.SubElement(normal_font, q('name'), {'val': 'Arial'})
    fonts.set('count', str(len(fonts)))
    for xf in styles.find(q('cellStyleXfs')):
        xf.set('fontId', str(len(fonts) - 1))
    # Keep explicit newlines visible in both shared-string and inline-string
    # cells. Conservative measured widths prevent unexpected second wrapping.
    for alignment in styles.iter(q('alignment')):
        alignment.set('wrapText', '1')
    files['xl/styles.xml'] = ET.tostring(styles, encoding='utf-8', xml_declaration=True)
    old_cols = sheet.find(q('cols'))
    if old_cols is not None:
        sheet.remove(old_cols)
    cols = ET.Element(q('cols'))
    sheet.insert(list(sheet).index(sheet.find(q('sheetData'))), cols)
    for index, points in enumerate(grid['widths'], 1):
        # Round cumulative pixel edges so a mixed-schema fine grid does not
        # accumulate rounding errors across many narrow physical columns.
        start = round(sum(grid['widths'][:index - 1]) * MASTER_COLUMN_SCALE * 96 / 72)
        end = round(sum(grid['widths'][:index]) * MASTER_COLUMN_SCALE * 96 / 72)
        width = math.floor((end - start) / 7 * 256) / 256
        ET.SubElement(cols, q('col'), {'min': str(index), 'max': str(index),
                                    'width': str(width), 'customWidth': '1'})
    last_column = grid['last_column']
    for name in ('sheetPr', 'printOptions', 'pageMargins', 'pageSetup', 'headerFooter', 'rowBreaks', 'colBreaks'):
        for elem in sheet.findall(q(name)):
            sheet.remove(elem)
    prop = ET.Element(q('sheetPr'))
    ET.SubElement(prop, q('pageSetUpPr'), {'fitToPage': '0', 'autoPageBreaks': '0'})
    sheet.insert(0, prop)
    ET.SubElement(sheet, q('printOptions'), {'horizontalCentered': '1', 'gridLines': '0'})
    ET.SubElement(sheet, q('pageMargins'), {'left': '0.3', 'right': '0.3', 'top': '0.4', 'bottom': '0.4', 'header': '0.15', 'footer': '0.15'})
    ET.SubElement(sheet, q('pageSetup'), {'paperSize': '9', 'orientation': 'portrait', 'scale': '100', 'fitToWidth': '1', 'fitToHeight': '0', 'useFirstPageNumber': '0'})
    breaks = ET.SubElement(sheet, q('rowBreaks'), {'count': str(page_count - 1), 'manualBreakCount': str(page_count - 1)})
    for i in range(1, page_count):
        ET.SubElement(breaks, q('brk'), {'id': str(i * MASTER_ROWS), 'min': '0', 'max': '16383', 'man': '1'})
    wb = ET.fromstring(files['xl/workbook.xml'])
    names = wb.find(q('definedNames'))
    if names is None:
        names = ET.Element(q('definedNames'))
        calc = wb.find(q('calcPr'))
        wb.insert(list(wb).index(calc) if calc is not None else len(wb), names)
    for n in list(names):
        if n.get('name') == '_xlnm.Print_Area':
            names.remove(n)
    ET.SubElement(names, q('definedName'), {'name': '_xlnm.Print_Area', 'localSheetId': '0'}).text = f"'单词总表'!$A$1:${last_column}${page_count * MASTER_ROWS}"
    files['xl/worksheets/sheet1.xml'] = ET.tostring(sheet, encoding='utf-8', xml_declaration=True)
    files['xl/workbook.xml'] = ET.tostring(wb, encoding='utf-8', xml_declaration=True)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)


def xlsx_fallback(layout, path):
    import xlsxwriter
    grid = master_grid(layout)
    area_columns, total_columns = grid['area_columns'], len(grid['widths'])
    wb = xlsxwriter.Workbook(path)
    sh = wb.add_worksheet('单词总表')
    sh.hide_gridlines(2)
    base = {'font_name': 'Arial Unicode MS', 'font_size': 9, 'valign': 'vcenter', 'text_wrap': True}
    normal = wb.add_format({**base, 'bottom': 1, 'bottom_color': '#DDDDDD'})
    header = wb.add_format({**base, 'bold': True, 'bg_color': '#F4F4F4', 'align': 'left'})
    groupfmt = wb.add_format({**base, 'bold': True, 'bg_color': '#E4E4E4'})
    titlefmt = wb.add_format({**base, 'bold': True, 'font_size': 12})
    for col, width in enumerate(grid['widths']):
        sh.set_column_pixels(col, col, width * 96 / 72)
    for row in range(len(layout) * MASTER_ROWS):
        sh.set_row(row, 6)
    def merged(r, c, h, w, text, fmt):
        sh.merge_range(r - 1, c, r + h - 2, c + w - 1, text, fmt)
    for page_idx, page in enumerate(layout):
        offset = page_idx * MASTER_ROWS
        merged(offset + 1, 0, 3, total_columns, '单词总表', titlefmt)
        for block in page:
            col, row = block['side'] * (area_columns + 1), offset + block['row']
            columns = block_columns(block, grid)
            merged(row, col, 3, area_columns, f"第{block['group']}组", groupfmt)
            row += 3
            for (left, span), name in zip(columns, block['headers']):
                merged(row, col + left, 3, span, name, header)
            row += 3
            for data in block['rows']:
                for (left, span), value in zip(columns, data['values']):
                    # strings_to_formulas disabled below by explicit write_string on anchor.
                    merged(row, col + left, data['units'], span, '', normal)
                    sh.write_string(row - 1, col + left, value, normal)
                row += data['units']
    wb.close()


def render_xlsx(layout, path, work):
    grid = master_grid(layout)
    requested = os.environ.get('VOCAB_XLSX_ENGINE', 'auto')
    if requested not in ('auto', 'xlsxwriter', 'artifact-tool'):
        raise ValueError('VOCAB_XLSX_ENGINE只能为auto、xlsxwriter或artifact-tool。')
    node = os.environ.get('VOCAB_NODE') or str(RUNTIME / 'node/bin/node')
    if not Path(node).is_file():
        node = shutil.which('node')
    module = os.environ.get('VOCAB_ARTIFACT_MODULE')
    package_dir = RUNTIME / 'node/node_modules/@oai/artifact-tool'
    if not module and package_dir.exists():
        package = json.loads((package_dir / 'package.json').read_text())
        exports = package.get('exports', {})
        entry = exports if isinstance(exports, str) else exports.get('.', {})
        if isinstance(entry, dict):
            entry = entry.get('import') or entry.get('default')
        if isinstance(entry, dict):
            entry = entry.get('default')
        entry = entry or package.get('module') or package.get('main')
        if entry:
            module = (package_dir / entry).resolve().as_uri()
    if requested == 'artifact-tool' and not (node and module):
        raise ValueError('找不到artifact-tool；可使用VOCAB_XLSX_ENGINE=xlsxwriter。')
    if requested != 'xlsxwriter' and node and module:
        payload = Path(work) / 'master-layout.json'
        payload.write_text(json.dumps({'pages': layout, 'rows_per_page': MASTER_ROWS,
                                       **grid}, ensure_ascii=False), encoding='utf-8')
        subprocess.run([node, str(ROOT / 'scripts/master.mjs'), str(payload), str(path), module], check=True)
        engine = 'artifact-tool'
    else:
        xlsx_fallback(layout, path)
        engine = 'xlsxwriter'
    patch_print_settings(path, len(layout), grid)
    return engine


def ocr_image(path):
    if not shutil.which('tesseract'):
        raise ValueError('未安装Tesseract；请由会话视觉读取图片，或安装本地OCR。')
    result = subprocess.run(['tesseract', str(path), 'stdout', '-l', 'eng+chi_sim', '--psm', '3'], capture_output=True, text=True)
    if result.returncode:
        raise ValueError('本地OCR失败：' + result.stderr[-400:])
    if not result.stdout.strip():
        raise ValueError('OCR未识别到文字；请由会话检查原图。')
    return result.stdout


class TextHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts, self.skip = [], 0
    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'):
            self.skip += 1
        if tag in ('p', 'div', 'br', 'li', 'tr', 'h1', 'h2'):
            self.parts.append('\n')
    def handle_endtag(self, tag):
        if tag in ('script', 'style'):
            self.skip = max(0, self.skip - 1)
    def handle_data(self, text):
        if not self.skip:
            self.parts.append(text)


def extract(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in ('.json', '.csv', '.tsv', '.txt', '.md'):
        return path.read_text(encoding='utf-8-sig')
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        lines = []
        for sh in wb:
            lines.append(f'\n[工作表：{sh.title}]')
            lines.extend('\t'.join(clean(v) for v in row) for row in sh.values if any(v is not None for v in row))
        wb.close()
        return '\n'.join(lines)
    if suffix == '.docx':
        with zipfile.ZipFile(path) as z:
            root = ET.fromstring(z.read('word/document.xml'))
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        result = []
        body = root.find('w:body', ns)
        for item in body:
            if item.tag.endswith('}p'):
                result.append(''.join(t.text or '' for t in item.findall('.//w:t', ns)))
            elif item.tag.endswith('}tbl'):
                for row in item.findall('w:tr', ns):
                    result.append('\t'.join(''.join(t.text or '' for t in cell.findall('.//w:t', ns)) for cell in row.findall('w:tc', ns)))
        return '\n'.join(result)
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        lines = []
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            if not text.strip():
                pdftoppm = shutil.which('pdftoppm') or str(RUNTIME / 'bin/override/pdftoppm')
                with tempfile.TemporaryDirectory() as tmp:
                    prefix = str(Path(tmp) / 'page')
                    subprocess.run([pdftoppm, '-f', str(i), '-l', str(i), '-singlefile', '-scale-to', '2400', '-png', str(path), prefix], check=True, capture_output=True)
                    text = ocr_image(prefix + '.png')
            lines.append(f'[第{i}页]\n{text}')
        return '\n\n'.join(lines)
    if suffix == '.epub':
        import posixpath
        with zipfile.ZipFile(path) as z:
            container = ET.fromstring(z.read('META-INF/container.xml'))
            opf_path = next(e.get('full-path') for e in container.iter() if e.tag.endswith('}rootfile'))
            opf = ET.fromstring(z.read(opf_path))
            items = {e.get('id'): e.get('href') for e in opf.iter() if e.tag.endswith('}item')}
            result = []
            for item in opf.iter():
                if item.tag.endswith('}itemref'):
                    html_path = posixpath.normpath(posixpath.join(posixpath.dirname(opf_path), items[item.get('idref')]))
                    parser = TextHTML()
                    parser.feed(z.read(html_path).decode('utf-8'))
                    result.append(''.join(parser.parts))
            return '\n\n'.join(result)
    if suffix in ('.png', '.jpg', '.jpeg', '.tif', '.tiff', '.bmp', '.webp', '.heic', '.heif'):
        from PIL import Image, ImageOps
        if suffix in ('.heic', '.heif'):
            try:
                from pillow_heif import register_heif_opener
            except ImportError as error:
                raise ValueError('HEIC/HEIF需要可选依赖pillow-heif；也可转成JPG或由会话直接看图。') from error
            register_heif_opener()
        with tempfile.TemporaryDirectory() as tmp:
            dst = Path(tmp) / 'image.png'
            with Image.open(path) as image:
                ImageOps.exif_transpose(image).convert('RGB').save(dst)
            return ocr_image(dst)
    raise ValueError('不支持此扩展名；DOC/XLS先另存为DOCX/XLSX，电子书可提供PDF页面或截图。')


def build(args):
    words, warnings = load_words(args.input)
    plan = schedule(words, args.daily)
    pages = pdf_layout(plan)
    master = master_layout(plan['groups'])
    out = Path(args.out).resolve()
    out.mkdir(parents=True, exist_ok=True)
    names = ('单词总表.xlsx', '艾宾浩斯复习打印册.pdf')
    if not args.overwrite and any((out / n).exists() for n in names):
        raise ValueError('输出文件已存在；请使用新目录，或明确加--overwrite。')
    # Keep staging on the same volume, including Windows D: outputs.
    with tempfile.TemporaryDirectory(prefix='.vocab-', dir=out) as temp:
        temp = Path(temp)
        render_pdf(pages, temp / names[1])
        engine = render_xlsx(master, temp / names[0], temp)
        for name in names:
            os.replace(temp / name, out / name)
    summary = {k: v for k, v in plan.items() if k not in ('groups', 'days')}
    summary.update({'pdf_pages': len(pages), 'xlsx_pages': len(master), 'xlsx_engine': engine, **warnings})
    if args.audit:
        audit = Path(args.audit)
        audit.parent.mkdir(parents=True, exist_ok=True)
        audit.write_text(json.dumps({'summary': summary, 'plan': plan, 'pdf': pages, 'master': master}, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(summary, ensure_ascii=False, indent=2))


def main():
    parser = argparse.ArgumentParser(description='艾宾浩斯式单词复习：本地排期与打印导出')
    sub = parser.add_subparsers(dest='command', required=True)
    b = sub.add_parser('build', help='从复核过的词表生成Excel和PDF')
    b.add_argument('input')
    b.add_argument('--daily', type=int, choices=[5, 10], required=True)
    b.add_argument('--out', required=True)
    b.add_argument('--audit', help='可选开发检查JSON路径')
    b.add_argument('--overwrite', action='store_true')
    e = sub.add_parser('extract', help='提取原始文字草稿，仍需核对原资料中的字段')
    e.add_argument('input')
    e.add_argument('--out', required=True)
    p = sub.add_parser('compare', help='在用户选择前计算5词和10词方案；不导出成品')
    p.add_argument('input')
    args = parser.parse_args()
    try:
        if args.command == 'build':
            build(args)
        elif args.command == 'compare':
            words, warnings = load_words(args.input)
            print(json.dumps({'plans': compare_plans(words), 'input_notes': warnings}, ensure_ascii=False, indent=2))
        else:
            target = Path(args.out)
            if target.exists():
                raise ValueError('提取目标已存在，请使用新文件名。')
            content = extract(args.input)
            if not content.strip():
                raise ValueError('未提取到文字，请查看页面图片。')
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(content, encoding='utf-8')
            print(f'已提取到 {target}；请对照原材料复核。')
    except (ValueError, OSError, subprocess.CalledProcessError) as error:
        parser.exit(2, f'无法完成：{error}\n')


if __name__ == '__main__':
    main()
