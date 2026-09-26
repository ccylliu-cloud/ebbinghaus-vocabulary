import collections
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from datetime import date
import zipfile
import xml.etree.ElementTree as ET
import vocab as v
from unittest.mock import patch


def words(n):
    return [{'english':f'word{i}', 'chinese':'示例词', 'ipa':'/wɜːd/', 'pos':'n.'} for i in range(n)]


class VocabularyTests(unittest.TestCase):
    def test_schedule_and_pagination_boundaries(self):
        for daily in (5,10):
            for n in (1,4,5,6,9,10,11,17,24,25,29,30,31,49,50,51,99,145,149,150,151,289,290,291,299,300,301,401):
                with self.subTest(daily=daily,n=n):
                    plan = v.schedule(words(n), daily, '2026-12-20')
                    groups = (n + daily - 1) // daily
                    self.assertEqual(plan['total_days'], groups + 29)
                    self.assertLessEqual(plan['peak_words'],daily*5)
                    self.assertEqual(sum(d['word_count'] for d in plan['days']),n*5)
                    seen = collections.defaultdict(list)
                    for day in plan['days']:
                        self.assertEqual((date.fromisoformat(day['date'])-date(2026,12,20)).days,day['day']-1)
                        for task in day['tasks']:
                            seen[task['group']].append((day['day'],task['visit']))
                    for group in range(1,groups+1):
                        self.assertEqual(seen[group],[(group+o,i+1) for i,o in enumerate((0,2,6,14,29))])
                    pages = v.pdf_layout(plan)
                    day_pages = collections.defaultdict(set)
                    actual = []
                    for p,page in enumerate(pages):
                        self.assertLessEqual(len(page),4)
                        for col in page:
                            day_pages[col['day']].add(p)
                            self.assertLessEqual(sum(b['height'] for b in col['blocks']),v.BODY_H)
                            for b in col['blocks']:
                                self.assertEqual(len(b['rows']),len(b['words']))
                                actual.append((col['day'],b['group'],b['visit']))
                    self.assertTrue(all(len(p)==1 for p in day_pages.values()))
                    expected=[(d['day'],t['group'],t['visit']) for d in plan['days'] for t in d['tasks']]
                    self.assertEqual(actual,expected)
                    for i,page in enumerate(pages[:-1]):
                        nextday=pages[i+1][0]['day']
                        nextcols=sum(c['day']==nextday for c in pages[i+1])
                        self.assertGreater(len(page)+nextcols,4)
                    master = v.master_layout(plan['groups'])
                    all_groups=[]
                    for page in master:
                        bounds={0:0,1:0}
                        for b in page:
                            self.assertGreaterEqual(b['row'],bounds[b['side']])
                            self.assertLessEqual(b['row']+b['units']-1,v.MASTER_ROWS)
                            bounds[b['side']]=b['row']+b['units']
                            all_groups.append(b['group'])
                    self.assertEqual(all_groups,list(range(1,groups+1)))

    def test_peak_and_early_days(self):
        for daily in (5,10):
            p=v.schedule(words(300),daily)
            self.assertEqual(p['peak_words'],daily*5)
            self.assertEqual(p['days'][29]['word_count'],daily*5)
            layout=v.pdf_layout(p)
            self.assertEqual([c['day'] for c in layout[0]],[1,2,3,4])
            peak=[c for page in layout for c in page if c['day']==30]
            self.assertGreater(len(peak),1)
            self.assertLessEqual(len(peak),4)
        small=v.schedule(words(1),5)
        self.assertEqual([c['day'] for p in v.pdf_layout(small) for c in p],[1,3,7,15,30])
        self.assertEqual(small['peak_words'],1)

    def test_wrapping_and_failure_is_explicit(self):
        data=words(300)
        data[0]['chinese']='把某物放回原来的位置'
        layout=v.pdf_layout(v.schedule(data,10))
        found=next(b for p in layout for c in p for b in c['blocks'] if b['group']==1)
        self.assertEqual(''.join(found['rows'][0]['lines']),data[0]['chinese'])
        data[0]['chinese']='长'*2000
        with self.assertRaises(ValueError):
            v.pdf_layout(v.schedule(data,10))
        with self.assertRaises(ValueError):
            v.master_layout(v.schedule(data,10)['groups'])

    def test_four_columns_stay_on_one_page(self):
        data=words(300)
        for index in list(range(270,280))+list(range(290,300)):
            data[index]['chinese']='需要换行的中文词义'
        pages=v.pdf_layout(v.schedule(data,10))
        peak_pages=[p for p in pages if any(c['day']==30 for c in p)]
        self.assertEqual(len(peak_pages),1)
        self.assertEqual([c['day'] for c in peak_pages[0]],[30,30,30,30])

    def test_long_prompt_uses_full_width_and_keeps_meaning(self):
        meaning = '（用于书信开头，向收信人表达礼貌及亲近关系）亲爱的朋友'
        data = [{'english': 'dear friend', 'chinese': meaning}]
        block = v.task_layout({'group': 1, 'visit': 1, 'words': data})
        row = block['rows'][0]
        self.assertTrue(row['stacked'])
        self.assertEqual(''.join(row['lines']), meaning)
        self.assertGreaterEqual(row['height'] - len(row['lines']) * 12, 24)
        short = v.task_layout({'group': 1, 'visit': 1, 'words': words(1)})
        self.assertFalse(short['rows'][0]['stacked'])

    def test_word_wrap_punctuation_and_measured_width(self):
        from reportlab.pdfbase.pdfmetrics import stringWidth
        lines = v.wrap('the Dragon Boat Festival', 70, 9)
        self.assertEqual(' '.join(lines), 'the Dragon Boat Festival')
        self.assertTrue(all(line == line.strip() for line in lines))
        source = '（用于书信开头，向收信人表示礼貌）亲爱的朋友；伙伴。'
        for width in (40, 55, 80, 100):
            lines = v.wrap(source, width, 9)
            self.assertEqual(''.join(lines), source)
            for line in lines:
                self.assertLessEqual(stringWidth(line, v.font_name(), 9), width)
                self.assertNotIn(line[0], '，。；）')
                self.assertNotEqual(line[-1], '（')

    def test_canonical_excel_columns_and_long_row_heights(self):
        data = words(40)
        for i, word in enumerate(data):
            word['chinese'] = '（用于书信开头，向收信人表达礼貌及亲近关系）亲爱的朋友'
            if i < 10: word.pop('ipa'); word.pop('pos')
            elif i < 20: word.pop('ipa')
            elif i < 30: word.pop('pos')
        for daily in (5, 10):
            layout = v.master_layout(v.schedule(data, daily)['groups'])
            grid = v.master_grid(layout)
            for page in layout:
                for block in page:
                    for row in block['rows']:
                        lines = max(len(value.split('\n')) for value in row['values'])
                        self.assertGreaterEqual(row['units'] * 6, lines * 12 + 8)
            with tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / 'master.xlsx'
                v.xlsx_fallback(layout, path)
                v.patch_print_settings(path, len(layout), grid)
                ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                with zipfile.ZipFile(path) as z:
                    sh = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
                    styles = ET.fromstring(z.read('xl/styles.xml'))
                columns = sh.find('m:cols', ns)
                self.assertEqual(len(columns), len(grid['widths']))
                self.assertLess(sum(float(c.get('width')) * 7 * .75 for c in columns), 510)
                self.assertEqual(sh.find('m:rowBreaks', ns).get('count'), str(len(layout) - 1))
                self.assertTrue(all(a.get('wrapText') == '1' for a in styles.findall('.//m:alignment', ns)))

    def test_input_validation_and_dedup(self):
        w=words(1)[0]
        result,info=v.normalize([w,w,{**w,'chinese':'另一释义'}])
        self.assertEqual(len(result),2)
        self.assertEqual(info['exact_duplicates_removed'],1)
        self.assertEqual(info['same_word_variants'],['word0'])
        for bad in ([],[{'english':'x'}],[{**w,'needs_review':True}],[{**w,'chinese':'坏\x01'}]):
            with self.assertRaises(ValueError): v.normalize(bad)
        minimal=v.normalize([{'english':'x','chinese':'叉'}])[0][0]
        self.assertEqual(set(minimal),{'english','chinese'})
        with self.assertRaises(ValueError): v.schedule(words(1),7)
        with self.assertRaises(ValueError): v.schedule(words(1),5,'2026-02-30')

    def test_artifact_pdf_answer_exclusion(self):
        from pypdf import PdfReader
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'test.pdf'
            data=[{'english':'SECRETANSWER','chinese':'苹果','ipa':'/SECRETIPA/','pos':'SECRET'}]
            v.render_pdf(v.pdf_layout(v.schedule(data,5,'2026-09-22')),path)
            pdf=PdfReader(path)
            text=''.join(p.extract_text() for p in pdf.pages)
            self.assertEqual(len(pdf.pages),2)
            self.assertEqual(text.count('苹果'),5)
            self.assertNotIn('SECRET',text)
            self.assertNotIn('掌握',text)
            self.assertNotIn('2026-',text)
            self.assertEqual(text.count('日期：____________'),5)
            for page in pdf.pages:
                self.assertAlmostEqual(float(page.mediabox.width),v.PDF_W,places=2)
                self.assertGreater(float(page.mediabox.width),float(page.mediabox.height))

    def test_optional_fields_and_compact_pairs(self):
        for daily in (5,10):
            for optional in ((),('pos',),('ipa',),('ipa','pos')):
                for n in (1,5,10,17,103,301):
                    with self.subTest(daily=daily,optional=optional,n=n):
                        raw=[{k:value for k,value in w.items() if k in ('english','chinese')+optional} for w in words(n)]
                        data,_=v.normalize(raw)
                        plan=v.schedule(data,daily)
                        pages=v.master_layout(plan['groups'])
                        seen=[]
                        for page in pages:
                            for block in page:
                                count=len(plan['groups'][block['group']-1]['words'])
                                seen.append(block['group'])
                                self.assertEqual(set(block['fields']),set(('english','chinese')+optional))
                                self.assertEqual(len(block['rows']),(count+1)//2 if not optional else count)
                                indexes=[i for row in block['rows'] for i in row['word_indexes'] if i is not None]
                                self.assertEqual(sorted(indexes),list(range(count)))
                                self.assertLessEqual(block['row']+block['units']-1,v.MASTER_ROWS)
                                if not optional:
                                    reading_order=[r['word_indexes'][col] for col in (0,1) for r in block['rows'] if r['word_indexes'][col] is not None]
                                    self.assertEqual(reading_order,list(range(count)))
                                    if count%2:
                                        self.assertEqual(block['rows'][-1]['values'][2:],['',''])
                        self.assertEqual(seen,list(range(1,len(plan['groups'])+1)))

    def test_mixed_groups_and_exact_source_information(self):
        raw=words(40)
        for i,w in enumerate(raw):
            if i<10: w.pop('ipa'); w.pop('pos')
            elif i<20: w.pop('ipa')
            elif i<30: w.pop('pos')
        raw[30]['ipa']='/ORIGINAL/'
        raw[31]['ipa']=''
        raw[32].pop('pos')
        normalized,_=v.normalize(raw)
        self.assertEqual(normalized[30]['ipa'],'/ORIGINAL/')
        self.assertNotIn('ipa',normalized[31])
        self.assertNotIn('pos',normalized[32])
        pages=v.master_layout(v.schedule(normalized,10)['groups'])
        blocks=[b for p in pages for b in p]
        self.assertEqual([b['headers'] for b in blocks],[['中文','英文','中文','英文'],['中文','英文','词性'],['中文','英文','音标'],['中文','英文','音标','词性']])
        grid=v.master_grid(pages)
        self.assertEqual(sum(grid['widths']),542)
        for b in blocks:
            spans=v.block_columns(b,grid)
            self.assertEqual(sum(span for _,span in spans),grid['area_columns'])
            self.assertTrue(all(span>0 for _,span in spans))
        self.assertEqual(blocks[3]['rows'][1]['values'][2],'')

    def test_compare_reports_actual_peak_before_selection(self):
        comparison=v.compare_plans(words(7))
        self.assertEqual(comparison,[{'daily':5,'word_count':7,'group_count':2,'total_days':31,'peak_words':5},{'daily':10,'word_count':7,'group_count':1,'total_days':30,'peak_words':7}])
        self.assertEqual([p['peak_words'] for p in v.compare_plans(words(300))],[25,50])

    def test_source_formats(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)
            for ext,text in (('csv','英文,中文,音标,词性\napple,苹果,/a/,n.\n'),
                             ('tsv','中文\t英文\t音标\t词性\n苹果\tapple\t/a/\tn.\n'),
                             ('json',json.dumps([{'english':'apple','chinese':'苹果','ipa':'/a/','pos':'n.'}]))):
                f=p/f'test.{ext}'; f.write_text(text,encoding='utf-8')
                self.assertEqual(v.load_words(f)[0][0]['chinese'],'苹果')
            # Minimal source fixture: creation here is for testing ingestion only.
            f=p/'test.xlsx'; wb=Workbook(); sh=wb.active
            sh.append(['英文','中文','音标','词性']); sh.append(['apple','苹果','/a/','n.']); wb.save(f)
            self.assertEqual(v.load_words(f)[0][0]['english'],'apple')
            self.assertIn('苹果',v.extract(f))
            f=p/'test.docx'
            with zipfile.ZipFile(f,'w') as z:
                z.writestr('word/document.xml','<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>apple 苹果</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r><w:t>banana</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>香蕉</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:body></w:document>')
            self.assertIn('apple 苹果',v.extract(f))
            self.assertIn('banana\t香蕉',v.extract(f))

    def test_standalone_cli_exports_and_protects_existing_files(self):
        from openpyxl import load_workbook
        from pypdf import PdfReader
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            # Spaces and non-ASCII paths are common in parents' folders.
            source = root / '已核对 词表.json'
            source.write_text(json.dumps([{'english': '=1+1', 'chinese': '原文符号'},
                                          {'english': 'apple', 'chinese': '苹果'}]), encoding='utf-8')
            target = root / '打印 成品'
            env = {**os.environ, 'VOCAB_XLSX_ENGINE': 'xlsxwriter', 'PYTHONUTF8': '1'}
            command = [sys.executable, str(v.ROOT / 'scripts/vocab.py'), 'build', str(source),
                       '--daily', '5', '--out', str(target)]
            first = subprocess.run(command, capture_output=True, text=True, encoding='utf-8', env=env)
            self.assertEqual(first.returncode, 0, first.stderr)
            summary = json.loads(first.stdout)
            self.assertEqual(summary['xlsx_engine'], 'xlsxwriter')
            self.assertEqual(summary['total_days'], 30)
            self.assertEqual({p.name for p in target.iterdir()}, {'单词总表.xlsx', '艾宾浩斯复习打印册.pdf'})
            excel = target / '单词总表.xlsx'
            before = excel.read_bytes()
            wb = load_workbook(excel)
            sheet = wb.active
            cells = [c for row in sheet for c in row if c.value is not None]
            original = next(c for c in cells if c.value == '=1+1')
            self.assertEqual(original.data_type, 's')
            self.assertNotIn('音标', [c.value for c in cells])
            self.assertNotIn('词性', [c.value for c in cells])
            self.assertEqual(sheet.page_setup.orientation, 'portrait')
            self.assertEqual(str(sheet.page_setup.paperSize), str(sheet.PAPERSIZE_A4))
            self.assertTrue(sheet.print_area)
            wb.close()
            pdf = PdfReader(target / '艾宾浩斯复习打印册.pdf')
            text = ''.join(p.extract_text() for p in pdf.pages)
            self.assertNotIn('apple', text)
            self.assertNotIn('=1+1', text)
            again = subprocess.run(command, capture_output=True, env=env)
            self.assertNotEqual(again.returncode, 0)
            self.assertEqual(excel.read_bytes(), before)

    def test_explicit_engine_rejects_typo(self):
        with patch.dict(os.environ, {'VOCAB_XLSX_ENGINE': 'typo'}):
            with self.assertRaises(ValueError):
                v.render_xlsx(v.master_layout(v.schedule(words(5), 5)['groups']), 'unused.xlsx', '.')


if __name__=='__main__':
    unittest.main(verbosity=2)
